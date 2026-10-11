"""PowerPoint extraction with slide provenance and optional visual analysis."""

import subprocess
import tempfile
from pathlib import Path
from typing import Any, Iterator

from pptx import Presentation

from backend.core.config import settings
from backend.core.llm import call_llm
from backend.core.logging import logger
from services.ingestion.prompts import SLIDE_PROMPT


def _render_slides(path: Path, output_dir: Path) -> bool:
    """Render a deck through LibreOffice when available."""
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
        pdf_path = output_dir / f"{path.stem}.pdf"
        if not pdf_path.exists():
            return False
        import fitz

        document = fitz.open(str(pdf_path))
        for index, page in enumerate(document):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(
                str(output_dir / f"slide_{index + 1}.png")
            )
        document.close()
        return True
    except Exception as exc:
        logger.warning("Slide rendering failed for %s: %s", path, exc)
        return False


def extract_text_and_images(
    pptx_path: str | Path, course_id: str, source_id: str
) -> Iterator[dict[str, Any]]:
    """Yield one chunk per slide, using visual analysis when rendering succeeds."""
    path = Path(pptx_path)
    output_dir = Path(settings.DATA_DIR) / "processed" / source_id
    output_dir.mkdir(parents=True, exist_ok=True)
    rendered = _render_slides(path, output_dir)
    presentation = Presentation(str(path))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        notes = ""
        if slide.has_notes_slide:
            notes = "\n".join(
                shape.text.strip()
                for shape in slide.notes_slide.notes_text_frame.paragraphs
                if shape.text.strip()
            )
        visible_text = "\n".join(texts)
        combined = f"{visible_text}\n{notes}".strip()
        visual: dict[str, Any] = {"visible_text": visible_text, "speaker_notes": notes}
        if rendered:
            image_path = output_dir / f"slide_{slide_number}.png"
            try:
                response = call_llm("A", f"{SLIDE_PROMPT}\n{combined}", images=[image_path])
                if isinstance(response.data, dict):
                    visual.update(response.data)
            except Exception as exc:
                visual["warning"] = str(exc)
        yield {
            "course_id": course_id,
            "source_id": source_id,
            "text": combined or "Slide with no extractable text",
            "visual": visual,
            "slide": slide_number,
            "modality": "text",
            "extraction_confidence": 0.9 if rendered else 0.75,
        }
